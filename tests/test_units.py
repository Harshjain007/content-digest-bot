"""Offline sanity checks for the pure logic — no network, no API keys.

Run:  python -m tests.test_units
"""
import json
import os
import tempfile

from content_digest_bot import store
from content_digest_bot.extractors import classify
from content_digest_bot.format_telegram import md_to_telegram_html, split_html
from content_digest_bot.github_api import is_github_url


def test_classify():
    assert classify("https://youtu.be/aircAruvnKk")[0] == "youtube"
    assert classify("https://www.youtube.com/watch?v=x")[0] == "youtube"
    assert classify("https://instagram.com/reel/ABC")[0] == "instagram"
    assert classify("https://example.com/post")[0] == "article"
    assert classify("what is RAG")[0] == "topic"
    # leftover text is preserved alongside the URL
    kind, url, rest = classify("https://example.com/p why does this matter")
    assert kind == "article" and rest == "why does this matter"


def test_is_github_url():
    assert is_github_url("https://github.com/psf/requests")
    assert not is_github_url("https://example.com")
    assert not is_github_url(None)


def test_markdown_to_html():
    html = md_to_telegram_html("## Title\n- **bold** item\n`code`")
    assert "<b>Title</b>" in html
    assert "• <b>bold</b> item" in html
    assert "<code>code</code>" in html
    # HTML in the source is escaped, not passed through to Telegram
    assert "&lt;script&gt;" in md_to_telegram_html("<script>alert(1)</script>")
    # links become anchors
    assert '<a href="https://x.com">t</a>' in md_to_telegram_html("[t](https://x.com)")


def test_split_html_balances_tags():
    md = "\n".join(f"<b>line {i} with some padding text</b>" for i in range(400))
    chunks = split_html(md, limit=500)
    assert len(chunks) > 1
    for c in chunks:
        assert len(c) <= 700  # limit + reopened/closed tags
        assert c.count("<b>") == c.count("</b>")


def test_dedup_and_store(tmpdir):
    store.DATA_DIR = tmpdir
    store.RESOURCES = os.path.join(tmpdir, "resources.json")
    store.LEARNINGS = os.path.join(tmpdir, "learnings.json")
    store.COMBINED = os.path.join(tmpdir, "data.json")
    # Keep the run entirely inside tmpdir: without these the test would
    # rewrite the real site and push it to gh-pages.
    store.SITE_HTML = os.path.join(tmpdir, "index.html")
    store.SITE_TEMPLATE = os.path.join(tmpdir, "missing-template.html")
    store._publish_to_pages = lambda *a, **k: None

    entry = {"title": "LangChain", "description": "framework for LLM apps",
             "links": {"github": "https://github.com/langchain-ai/langchain"}}
    added, _ = store.add_resource(dict(entry))
    assert added

    # exactly the same card again -> rejected on the link
    added, reason = store.add_resource(dict(entry))
    assert not added and "link" in reason

    # same title, no link at all -> rejected on the title
    added, reason = store.add_resource(
        {"title": "LangChain", "description": "unrelated wording entirely"})
    assert not added and "title" in reason

    # same repo under a different title -> still rejected
    added, _ = store.add_resource({"title": "LangChain Core",
                                   "description": "totally unrelated words here",
                                   "links": entry["links"]})
    assert not added

    # the same repo dressed up differently: other scheme, www, .git, trailing
    # slash, different casing. All one link, so all rejected.
    for variant in ("http://github.com/langchain-ai/langchain",
                    "https://www.github.com/langchain-ai/langchain/",
                    "https://github.com/langchain-ai/langchain.git",
                    "https://GitHub.com/langchain-ai/LangChain"):
        added, reason = store.add_resource({
            "title": f"variant {variant}",
            "description": "words that share nothing with the first entry",
            "links": {"github": variant}})
        assert not added, f"{variant} should have been caught as a duplicate"
        assert "link" in reason

    # a learning card pointing at an already-filed link is a duplicate too,
    # even though it lands in the other store
    added, _ = store.add_learning({
        "description": "a write-up about the same repository",
        "links": "https://github.com/langchain-ai/langchain",
        "takeAways": ["it chains things"]})
    assert not added

    # genuinely different entry -> accepted
    added, _ = store.add_resource({
        "title": "Whisper",
        "description": "speech recognition model transcribing audio offline",
        "links": {"github": "https://github.com/openai/whisper"}})
    assert added

    # a learning card with a bare-string link is normalized to a dict on save
    added, _ = store.add_learning({
        "description": "how to run deliberate practice sessions each week",
        "links": "https://example.com/deep-work",
        "takeAways": ["block the calendar"]})
    assert added
    with open(store.LEARNINGS) as f:
        assert json.load(f)[0]["links"] == {"article": "https://example.com/deep-work"}

    # combined data.json is regenerated for the site
    with open(store.COMBINED) as f:
        assert len(json.load(f)) == 3


def test_dedupe_cleans_existing(tmpdir):
    """dedupe() collapses duplicates that predate the checks."""
    store.RESOURCES = os.path.join(tmpdir, "r2.json")
    store.LEARNINGS = os.path.join(tmpdir, "l2.json")
    store._save(store.RESOURCES, [
        {"title": "Flask", "description": "a micro web framework",
         "links": {"github": "https://github.com/pallets/flask"}},
        {"title": "Flask (again)", "description": "different words entirely",
         "links": {"github": "https://github.com/pallets/flask/"}},
        {"title": "Django", "description": "batteries included web framework",
         "links": {"github": "https://github.com/django/django"}},
    ])
    store._save(store.LEARNINGS, [])
    dropped = store.dedupe()
    assert len(dropped) == 1
    with open(store.RESOURCES) as f:
        kept = json.load(f)
    assert [e["title"] for e in kept] == ["Flask", "Django"]


def test_to_markdown_preserves_structure():
    """Documents reach the model as Markdown, not flattened text.

    Heading levels and tables are the structure the model actually uses; the
    old per-format parsing threw both away.
    """
    import io
    import docx
    from content_digest_bot.extractors import to_markdown

    d = docx.Document()
    d.add_heading("Quarterly Review", 1)
    d.add_heading("Revenue", 2)
    d.add_paragraph("Revenue grew across all regions.")
    t = d.add_table(rows=2, cols=2)
    for r, row in enumerate([["Region", "Q1"], ["EMEA", "120"]]):
        for c, v in enumerate(row):
            t.rows[r].cells[c].text = v
    buf = io.BytesIO()
    d.save(buf)

    md, _ = to_markdown(buf.getvalue(), ".docx")
    assert "# Quarterly Review" in md          # h1 survives as h1
    assert "## Revenue" in md                  # h2 is distinguishable from h1
    assert "| Region | Q1 |" in md             # a real markdown table
    assert "| --- |" in md

    # the cap is honoured, so MAX_INPUT_CHARS actually bounds token spend
    short, _ = to_markdown(buf.getvalue(), ".docx", max_chars=20)
    assert len(short) == 20

    # unreadable input raises rather than sending noise to the model
    try:
        to_markdown(b"not a document", ".docx")
        raise AssertionError("expected a failure on garbage input")
    except AssertionError:
        raise
    except Exception:
        pass


def test_document_title_falls_back_to_heading():
    """A document card must not be titled with its own URL.

    MarkItDown carries no title: pdfminer exposes no metadata and the docx
    path goes through mammoth, which drops core properties. Without a fallback
    every paper filed by link was titled "https://arxiv.org/pdf/..." — which
    also collapses title-based de-duplication down to URL matching.
    """
    import io
    import docx
    from content_digest_bot.extractors import extract_document, _md_title

    d = docx.Document()
    d.core_properties.title = "Ignored By Mammoth"
    d.add_heading("Attention Is All You Need", 1)
    d.add_paragraph("We propose a new simple network architecture. " * 4)
    buf = io.BytesIO()
    d.save(buf)

    entry = extract_document(buf.getvalue(), ".docx",
                             url="https://example.com/paper.docx")
    assert entry["title"] == "Attention Is All You Need", entry["title"]
    assert entry["suffix"] == ".docx" and entry["is_doc"] and not entry["is_pdf"]

    # no heading at all -> fall back to the URL rather than inventing one
    assert _md_title("just body text, no heading") is None
    assert _md_title("## Second level still counts") == "Second level still counts"

    # unreadable bytes fail cleanly instead of raising at the call site
    bad = extract_document(b"not a document", ".docx", url="https://x.test/a")
    assert bad["failed"] and bad["text"] is None


def test_pdf_title_read_from_metadata():
    """A PDF is titled from its own /Title, not from its URL.

    MarkItDown's pdfminer backend exposes no metadata and a PDF's Markdown has
    no `#` heading to fall back on, so without this every paper filed by link
    was titled "https://arxiv.org/pdf/...".
    """
    from content_digest_bot.extractors import _pdf_title

    pdf = (b"%PDF-1.4\n"
           b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
           b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
           b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]>>endobj\n"
           b"4 0 obj<</Title(Attention Is All You Need)>>endobj\n"
           b"trailer<</Root 1 0 R/Info 4 0 R>>\n%%EOF\n")
    assert _pdf_title(pdf) == "Attention Is All You Need"
    # unreadable input must not raise — it is only ever a nice-to-have
    assert _pdf_title(b"not a pdf at all") is None


def test_doc_suffix_routing():
    """Link route and upload route accept the same set of formats."""
    from content_digest_bot.extractors import doc_suffix, DOC_SUFFIXES

    assert doc_suffix("https://x.test/deck.pptx") == ".pptx"
    assert doc_suffix("https://x.test/sheet.XLSX") == ".xlsx"
    assert doc_suffix("https://x.test/paper.pdf?download=1") == ".pdf"
    assert doc_suffix("https://x.test/post") is None
    # arXiv links carry no extension at all — they are routed by what the
    # server says it served, which is why the content-type arm has to exist.
    assert doc_suffix("https://arxiv.org/pdf/2309.06180") is None
    assert doc_suffix("https://arxiv.org/pdf/2309.06180",
                      "application/pdf") == ".pdf"
    assert ".pptx" in DOC_SUFFIXES and ".xlsx" in DOC_SUFFIXES


def test_error_messages_do_not_leak():
    """Only messages we wrote ourselves reach the user.

    Raw exception text routinely carries local paths and request URLs, and it
    used to be interpolated straight into the Telegram reply.
    """
    from content_digest_bot.errors import BotError, user_message

    assert user_message(BotError("Ollama isn't reachable.")) == "Ollama isn't reachable."

    leaky = FileNotFoundError("[Errno 2] No such file: '/Users/harshjain/.env'")
    out = user_message(leaky)
    assert "/Users/harshjain" not in out and ".env" not in out
    assert out == "Something went wrong on my side. It's in the log."

    # an empty curated message must not produce a bare "❌ "
    assert user_message(BotError("  ")) != ""


def test_synthesize_falls_back_when_ollama_is_down():
    """A sleeping local daemon must not cost the user their submission."""
    from content_digest_bot import synthesize as sy
    from content_digest_bot.errors import BotError

    real_ollama, real_anthropic = sy._synthesize_ollama, sy._synthesize_anthropic
    real_provider, real_key = sy.LLM_PROVIDER, sy.ANTHROPIC_API_KEY
    try:
        sy.LLM_PROVIDER = "ollama"
        sy.ANTHROPIC_API_KEY = "sk-test"
        sy._synthesize_ollama = lambda *a, **k: (_ for _ in ()).throw(
            BotError("Ollama isn't reachable."))
        sy._synthesize_anthropic = lambda *a, **k: "card from anthropic"

        assert sy._call("prompt") == "card from anthropic"
        assert sy.used_fallback.get() is True   # so the bot can say so

        # with no key there is nothing to fall back to: surface the real cause
        sy.ANTHROPIC_API_KEY = ""
        try:
            sy._call("prompt")
            raise AssertionError("expected the Ollama error to propagate")
        except BotError as e:
            assert "reachable" in str(e)

        # the happy path leaves the flag clear, so no spurious note appears
        sy.ANTHROPIC_API_KEY = "sk-test"
        sy._synthesize_ollama = lambda *a, **k: "card from ollama"
        assert sy._call("prompt") == "card from ollama"
        assert sy.used_fallback.get() is False
    finally:
        sy._synthesize_ollama, sy._synthesize_anthropic = real_ollama, real_anthropic
        sy.LLM_PROVIDER, sy.ANTHROPIC_API_KEY = real_provider, real_key


def test_format_card_is_safe_and_readable():
    """Cards replaced a raw JSON dump; entries are model-written, so escape."""
    from content_digest_bot.format_telegram import format_card

    entry = {
        "title": "Tool <script>alert(1)</script> & co",
        "description": "Does <b>things</b>",
        "problem": "Solves 5 > 3 problems",
        "how this works": "1. Install it 2. Configure it 3. Run it",
        "links": {"github": "https://github.com/a/b",
                  "website": "javascript:alert(1)",
                  "article": ""},
    }
    out = format_card(entry, "tool", saved=True, register_url="https://x.test/r")

    # model text is escaped, so it cannot break Telegram's HTML parser
    assert "<script>" not in out and "&lt;script&gt;" in out
    assert "&amp; co" in out and "5 &gt; 3" in out
    # only http(s) links survive
    assert '<a href="https://github.com/a/b">GitHub</a>' in out
    assert "javascript:" not in out
    # a single-line numbered run becomes separate steps
    assert out.count("2. Configure it") == 1
    assert "\n　1. Install it" in out
    # the register link rides along instead of a separate message
    assert '<a href="https://x.test/r">Register</a>' in out

    # a duplicate says why, without pretending it saved
    dup = format_card(entry, "tool", saved=False, reason="already filed under this link")
    assert "Already filed" in dup and "already filed under this link" in dup


def test_keywords_handles_non_string_fields():
    """takeAways is a list, and every comparison funnels through _keywords.

    This raised `AttributeError: 'list' object has no attribute 'lower'` in
    production, which meant learning cards never got de-duplicated at all —
    add_learning crashed instead of returning a verdict.
    """
    from content_digest_bot.store import _keywords, _is_duplicate

    assert _keywords(["use Ditto", "try Vibe Coding"]) >= {"ditto", "vibe"}
    for weird in (None, 42, {"a": 1}, ["x", None, 7], ("t",)):
        _keywords(weird)              # must not raise

    entry = {"description": "free AI resources",
             "takeAways": ["use Ditto", "try Vibe Coding"]}
    assert _is_duplicate(entry, [dict(entry)]) is True


def test_model_calls_do_not_block_the_event_loop():
    """A long generation must not freeze every other chat.

    Model calls are synchronous and ran directly on the loop, so a 4096-token
    local generation stalled polling, /help, and even editing the status
    message the user was staring at.
    """
    import asyncio
    import time
    from content_digest_bot.bot import _model_call, _provider_note
    from content_digest_bot.synthesize import used_fallback

    def slow():
        time.sleep(0.6)
        return "generated"

    async def other_traffic():
        ticks = 0
        for _ in range(6):
            await asyncio.sleep(0.1)
            ticks += 1
        return ticks

    async def scenario():
        t0 = time.time()
        out, ticks = await asyncio.gather(_model_call(slow), other_traffic())
        return out, ticks, time.time() - t0

    out, ticks, elapsed = asyncio.run(scenario())
    assert out == "generated"
    assert ticks == 6, "the loop was blocked while the model ran"
    assert elapsed < 1.1, f"calls ran serially ({elapsed:.2f}s), not concurrently"

    # the fallback flag is set inside the worker thread, which gets a *copy* of
    # the context; it has to be carried back or the provider note goes missing.
    async def with_fallback():
        used_fallback.set(False)
        await _model_call(lambda: used_fallback.set(True) or "card")
        return used_fallback.get(), _provider_note()

    flag, note = asyncio.run(with_fallback())
    assert flag is True and "Anthropic" in note

    async def without_fallback():
        used_fallback.set(False)
        await _model_call(lambda: "card")
        return _provider_note()

    assert asyncio.run(without_fallback()) == ""


def main():
    with tempfile.TemporaryDirectory() as tmp:
        for name, fn in sorted(globals().items()):
            if not name.startswith("test_") or not callable(fn):
                continue
            fn(tmp) if fn.__code__.co_argcount else fn()
            print(f"  ok  {name}")
    print("\nRESULT: PASS")


if __name__ == "__main__":
    main()
